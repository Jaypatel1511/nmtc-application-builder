"""The Key Metrics row states what its figure is a share of, on every surface.

THE DEFECT (1.7.1 R9)
=====================

``excel_builder.py`` has rendered

    "Deep/Severe Distress Concentration " + Q25_QEI_BASIS_SUFFIX_SHEET

on the Summary Dashboard since 1.3.0 S4, when that cell was found carrying a
raw float under a percent format with no denominator in its label at all.
``word_builder`` and ``pdf_builder`` render the same metric, in the Key
Metrics table on the first page of the Executive Summary, and through 1.7.0
their label was

    "Deep/Severe Distress Concentration"

— no denominator, not even the word QEI. Question 25's two commitments are
measured on **QLICIs**; every distress share this package computes is
denominated in **QEI**. A CDE copying this cell into Question 25 files a QEI
figure against a QLICI commitment, which is the same defect 1.3.0 S4 removed
from the workbook, left live on the two documents a reviewer actually reads.

So the workbook said what the figure was a share of and Word and PDF did not:
the surface-drift shape this package has shipped a blocking defect from twice
(1.6.2's Q25 note, 1.7.0's docs sample) — a fix lands on one surface and the
others keep the old text.

WHY THE BARE CLAUSE AND NOT EITHER POINTER SUFFIX
=================================================

``_question_25`` offers three forms, and the Excel twin's SHAPE (label, space,
parenthesised denominator statement) is what is matched, not its wording:

  * ``Q25_QEI_BASIS_SUFFIX_SHEET`` names the workbook's ``Q25 Basis Note``
    tab. No flowing document has a tab; on Word or PDF it would point a
    federal reviewer at a sheet that is not in the file they are holding.
  * ``Q25_QEI_BASIS_SUFFIX`` says "see the basis note below", and its own note
    in ``_question_25`` justifies that pointer by locality — the note "renders
    as the last row of the same two-column table, in the same column of the
    same page, a few lines under the figure". In the Key Metrics table it is
    twenty pages away, in Section B, so the justification does not hold here.
  * ``Q25_QEI_BASIS_CLAUSE`` — the denominator with no pointer attached — is
    what 1.7.1 R8 already put in the Executive Summary sentence three lines
    above this table. Using it here makes the sentence and the row state the
    denominator identically on the same page, rather than introducing a third
    phrasing of one fact into one section.

WHAT THIS GATE ASSERTS, AND HOW IT IS BOUNDED
=============================================

**A gate that searches a whole rendered document is vacuous here**, because
the clause legitimately appears three times in each document: the Executive
Summary sentence (R8), this row, and the Question 25 basis note in Section B.
A document-wide ``assert clause in text`` would stay green with this row's
label reverted to its 1.7.0 spelling. So every assertion below is made against
a BOUNDED REGION — one Word table cell, one workbook cell, and, for the PDF,
the extracted text between this row's label and the next metric's label — and
``test_the_region_excludes_the_other_occurrences`` proves the boundary holds
by asserting the neighbouring occurrences are OUTSIDE it.

Two-sided: the clause is on the label AND the row still carries its figure.
"""
from __future__ import annotations

import re

import pytest

from nmtcapp.renderers._question_25 import (
    Q25_QEI_BASIS_CLAUSE, Q25_QEI_BASIS_SUFFIX_SHEET,
)

#: The stem all three surfaces share. Typed once here, on purpose: this module
#: is what compares the three renderers' own typed copies against each other.
LABEL_STEM = "Deep/Severe Distress Concentration"

#: The metric rendered on the row AFTER this one, on all three surfaces. It is
#: the PDF region's right-hand boundary, so it is named rather than inlined.
NEXT_LABEL = "NMTC Eligibility Rate"


@pytest.fixture(scope="module")
def rendered(tmp_path_factory) -> dict:
    from nmtcapp.core.application import Application
    from tests.test_rendered_output_baseline import (
        APPLICATION_ROUND, REQUESTED_ALLOCATION, _cde, _pipeline,
    )

    app = Application(cde=_cde(), requested_allocation=REQUESTED_ALLOCATION,
                      application_round=APPLICATION_ROUND)
    app.add_pipeline(_pipeline())
    out = str(tmp_path_factory.mktemp("keymetrics"))
    paths = app.generate(out, formats=["word", "pdf", "excel"])
    assert {"word", "pdf", "excel"} <= set(paths), (
        f"the fixture rendered {sorted(paths)}; this gate needs all three"
    )
    share = app.analyze().pipeline_result.distress_breakdown["pct_deep_or_severe"]
    return {"paths": paths, "share": share}


# ---------------------------------------------------------------------------
# The bounded region, one extractor per surface
# ---------------------------------------------------------------------------

def _word_region(path: str) -> str:
    """The ONE table cell carrying this label, and nothing else in the file."""
    import docx

    cells = [
        row.cells[0].text
        for table in docx.Document(path).tables
        for row in table.rows
        if row.cells and row.cells[0].text.strip().startswith(LABEL_STEM)
    ]
    assert len(cells) == 1, (
        f"expected exactly one Word table cell starting {LABEL_STEM!r}; "
        f"found {len(cells)}: {cells!r}. The extractor is no longer reading "
        "one row, so nothing below is bounded."
    )
    return cells[0]


def _pdf_region(path: str) -> str:
    """The extracted text from this row's label up to the next metric's.

    The label wraps inside its cell, so the region is flattened to one line
    before it is searched — the same normalisation
    tests/test_headline_distress_denominator.py applies for the same reason.
    """
    from pypdf import PdfReader

    text = "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    start = text.index(LABEL_STEM)
    end = text.index(NEXT_LABEL, start)
    region = re.sub(r"\s+", " ", text[start:end]).strip()
    assert len(region) < 200, (
        f"the PDF region is {len(region)} characters: {region!r}. It is "
        "supposed to be one table row. A region this large is how a gate "
        "passes on text 120 lines from the site it claims to check."
    )
    return region


def _excel_region(path: str) -> str:
    """The ONE workbook cell carrying this label."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    values = [
        cell.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith(LABEL_STEM)
    ]
    assert len(values) == 1, (
        f"expected exactly one workbook cell starting {LABEL_STEM!r}; found "
        f"{len(values)}: {values!r}"
    )
    return values[0]


_REGIONS = {"word": _word_region, "pdf": _pdf_region, "excel": _excel_region}


# ---------------------------------------------------------------------------
# The assertions
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("fmt", ["word", "pdf", "excel"])
def test_the_key_metrics_row_states_its_denominator(rendered, fmt):
    """Bounded to the row: the label says what the figure is a share of."""
    region = _REGIONS[fmt](rendered["paths"][fmt])
    assert Q25_QEI_BASIS_CLAUSE in region, (
        f"{fmt}: the Key Metrics row labelled {LABEL_STEM!r} states no "
        f"denominator. Expected {Q25_QEI_BASIS_CLAUSE!r} on the label — the "
        "figure is a share of QEI and Question 25's commitments are measured "
        f"on QLICIs.\n\nRegion read: {region!r}"
    )


@pytest.mark.parametrize("fmt", ["word", "pdf"])
def test_the_row_still_carries_its_figure(rendered, fmt):
    """THE OTHER SIDE. A label with no number beside it is not a fix.

    The workbook stores this cell as a float under a percent format, so the
    figure is not text there and this half is asserted on the two flowing
    surfaces, where the value is rendered text in the same row.
    """
    share = f"{rendered['share']:.0%}"
    if fmt == "word":
        import docx

        rows = [
            row for table in docx.Document(rendered["paths"][fmt]).tables
            for row in table.rows
            if row.cells and row.cells[0].text.strip().startswith(LABEL_STEM)
        ]
        value = rows[0].cells[1].text.strip()
    else:
        value = _pdf_region(rendered["paths"][fmt])
    assert share in value, (
        f"{fmt}: the Key Metrics row carries the denominator but not the "
        f"figure {share!r} it qualifies. Read: {value!r}"
    )


@pytest.mark.parametrize("fmt", ["word", "pdf"])
def test_the_region_excludes_the_other_occurrences(rendered, fmt):
    """THE BOUNDARY IS THE WHOLE GATE. Prove it is not the whole document.

    ``Q25_QEI_BASIS_CLAUSE`` renders at least three times in each of these
    documents: the R8 Executive Summary sentence, this row, and the Question
    25 basis note in Section B. A document-wide assertion would therefore stay
    green with this row reverted, which is exactly the shape a sibling
    implementation shipped. This asserts the region contains ONE of the three
    — so removing that one reddens the gate above.
    """
    path = rendered["paths"][fmt]
    region = _REGIONS[fmt](path)

    if fmt == "word":
        import docx

        whole = "\n".join(p.text for p in docx.Document(path).paragraphs)
        whole += "\n" + "\n".join(
            cell.text for table in docx.Document(path).tables
            for row in table.rows for cell in row.cells
        )
    else:
        from pypdf import PdfReader

        whole = "\n".join((pg.extract_text() or "")
                          for pg in PdfReader(path).pages)
    whole_flat = re.sub(r"\s+", " ", whole)

    assert whole_flat.count(Q25_QEI_BASIS_CLAUSE) >= 3, (
        f"{fmt}: the clause renders {whole_flat.count(Q25_QEI_BASIS_CLAUSE)} "
        "time(s) in the whole document. This proof needs the neighbouring "
        "occurrences to exist — if they were removed, say so and re-derive "
        "this count rather than lowering it."
    )
    assert region.count(Q25_QEI_BASIS_CLAUSE) == 1, (
        f"{fmt}: the bounded region holds "
        f"{region.count(Q25_QEI_BASIS_CLAUSE)} occurrences of the clause. It "
        "is supposed to hold exactly the one on this row; anything more means "
        "the boundary has slipped and the gate above can pass on a "
        f"neighbour's text.\n\nRegion read: {region!r}"
    )
    assert "committed to deep" not in region.lower(), (
        f"{fmt}: the Executive Summary sentence has leaked into the region "
        f"this gate calls one table row: {region!r}"
    )


def test_the_three_surfaces_share_one_label_stem(rendered):
    """They cannot drift into three different names for one metric."""
    stems = {}
    for fmt in ("word", "pdf", "excel"):
        region = _REGIONS[fmt](rendered["paths"][fmt])
        assert region.startswith(LABEL_STEM), (
            f"{fmt}: the Key Metrics label no longer starts {LABEL_STEM!r}; "
            f"it reads {region[:80]!r}. If the metric was renamed, rename it "
            "on all three surfaces and here — do not delete the check."
        )
        stems[fmt] = region[:len(LABEL_STEM)]
    assert len(set(stems.values())) == 1, f"the label stem drifted: {stems}"


def test_the_workbooks_wording_is_still_its_own_and_is_not_copied_here():
    """The shape is shared; the pointer is not.

    The workbook points at a TAB and the flowing documents have none. If a
    later round "unifies" these labels by giving Word and PDF the sheet
    suffix, the two filing documents start naming a sheet they do not
    contain — so the two forms are asserted to stay distinct.
    """
    import inspect

    from nmtcapp.renderers import pdf_builder, word_builder

    for module in (word_builder, pdf_builder):
        # Comment lines are stripped: word_builder's own note NAMES the
        # workbook suffix to explain why it is not used, and a scanner that
        # cannot tell a citation from a call would fail on the explanation.
        source = "\n".join(
            line for line in inspect.getsource(module).splitlines()
            if not line.lstrip().startswith("#")
        )
        assert "Q25_QEI_BASIS_SUFFIX_SHEET" not in source, (
            f"{module.__name__} references the WORKBOOK's basis suffix, which "
            f"reads {Q25_QEI_BASIS_SUFFIX_SHEET!r} and names a tab this "
            "document does not have."
        )
    # Retyping the clause instead of reading the constant is adjudicated by
    # tests/test_headline_distress_denominator.py, which walks the installed
    # package rather than two named modules. Not repeated here.
